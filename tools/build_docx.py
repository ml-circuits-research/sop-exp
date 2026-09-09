#!/usr/bin/env python3
"""Build the programmer guide from docs/GHID.md. Requires python-docx only."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/'docs/GHID.md'
OUTPUT=ROOT/'docs/SOP_Ghid_Experiment.docx'

def shade(element, color):
    pr=element.get_or_add_tcPr() if element.tag==qn('w:tc') else element.get_or_add_pPr()
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),color);pr.append(sh)

def inline(p,text):
    for part in re.split(r'(\*\*.*?\*\*|`[^`]+`)',text):
        if not part:continue
        r=p.add_run(part[2:-2] if part.startswith('**') else part[1:-1] if part.startswith('`') else part)
        if part.startswith('**'):r.bold=True
        if part.startswith('`'):
            r.font.name='DejaVu Sans Mono';r.font.size=Pt(9.0)
            r.font.color.rgb=RGBColor.from_string('293A4D')


def build():
    d=Document()
    sec=d.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.7);sec.bottom_margin=Cm(1.55)
    sec.left_margin=Cm(1.9);sec.right_margin=Cm(1.9)
    sec.header_distance=Cm(.65);sec.footer_distance=Cm(.65)
    styles=d.styles
    st=styles['Normal'];st.font.name='Calibri';st.font.size=Pt(10.5)
    st.font.color.rgb=RGBColor.from_string('192432')
    st.paragraph_format.space_after=Pt(6)
    st.paragraph_format.line_spacing=1.04
    st.paragraph_format.widow_control=True
    for name,size,color in [('Title',22,'12273E'),('Heading 1',17,'12273E'),('Heading 2',11.5,'175E6F')]:
        st=styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string(color)
        st.font.bold=name!='Title';st.paragraph_format.space_before=Pt(9 if name=='Heading 2' else 0)
        st.paragraph_format.space_after=Pt(7);st.paragraph_format.keep_with_next=True
    if 'Code Block' not in styles:
        st=styles.add_style('Code Block',1)
    st=styles['Code Block'];st.font.name='DejaVu Sans Mono';st.font.size=Pt(8.5)
    st.font.color.rgb=RGBColor.from_string('192432')
    st.paragraph_format.space_before=Pt(4);st.paragraph_format.space_after=Pt(7)
    st.paragraph_format.line_spacing=1.0;st.paragraph_format.left_indent=Cm(.2);st.paragraph_format.right_indent=Cm(.1)
    st.paragraph_format.keep_together=True
    header=sec.header.paragraphs[0];header.text='SOP  /  EXPERIMENTE ȘI VERIFICARE'
    header.style=styles['Normal'];header.paragraph_format.space_after=Pt(0)
    for r in header.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('617587')
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=footer.add_run('Ghid de lucru  ·  ');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('617587')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    lines=SOURCE.read_text().splitlines();i=0;first=True
    while i<len(lines):
        line=lines[i]
        if not line.strip():i+=1;continue
        if line=='<!-- page -->':d.add_page_break();i+=1;continue
        if line.startswith('```'):
            i+=1;code=[]
            while i<len(lines) and not lines[i].startswith('```'):code.append(lines[i]);i+=1
            p=d.add_paragraph('\n'.join(code),'Code Block');shade(p._p,'F2F5F8');i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].startswith('|'):
                cells=[x.strip() for x in lines[i].strip('|').split('|')]
                if not all(re.match(r'^:?-+:?$',x) for x in cells):rows.append(cells)
                i+=1
            table=d.add_table(rows=0,cols=len(rows[0]));table.autofit=False
            for j,row in enumerate(rows):
                cells=table.add_row().cells
                for k,text in enumerate(row):
                    cells[k].width=Cm(5.35 if k==0 else 11.85)
                    inline(cells[k].paragraphs[0],text)
                    for p in cells[k].paragraphs:
                        p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(4)
                        for r in p.runs:r.font.size=Pt(9.5)
                    if j==0:
                        shade(cells[k]._tc,'E6EEF3')
                        for r in cells[k].paragraphs[0].runs:r.bold=True
                    else:shade(cells[k]._tc,'F7F9FA')
                trPr=table.rows[-1]._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');trPr.append(no)
                if j==0:repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
            d.add_paragraph().paragraph_format.space_after=Pt(0)
            continue
        if line.startswith('# '):
            d.add_paragraph(line[2:], 'Title' if first else 'Heading 1');first=False;i+=1;continue
        if line.startswith('## '):d.add_paragraph(line[3:],'Heading 2');i+=1;continue
        para=[line];i+=1
        while i<len(lines) and lines[i].strip() and not lines[i].startswith(('#','```','|','<!--')):para.append(lines[i]);i+=1
        p=d.add_paragraph();inline(p,' '.join(para))
    d.core_properties.title='SOP: ce face experimentul și cum îi verificăm rezultatele'
    d.core_properties.subject='Ghid de lucru cu rezultate, contraexemple și criterii experimentale'
    d.core_properties.author=''
    d.core_properties.keywords='SOP, experimente, programe, verificare, limite'
    d.save(OUTPUT)
    print(OUTPUT)
if __name__=='__main__':build()
